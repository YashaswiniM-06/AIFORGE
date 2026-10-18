"""Bounded document parsing. Uploaded content is data; macros and links are never run."""
import sys, json, re, csv, io, hashlib, zipfile, logging
from pathlib import Path
from xml.sax.saxutils import escape
logging.disable(logging.CRITICAL)
MAX_PAGES = 12

def scan_codes(image):
    """Return type/location only. Payloads are deliberately never stored or followed."""
    try:
        import zxingcpp
        width,height=image.size; found=[]
        for result in zxingcpp.read_barcodes(image):
            if not result.valid: continue
            points=[result.position.top_left,result.position.top_right,result.position.bottom_left,result.position.bottom_right]
            xs=[point.x for point in points]; ys=[point.y for point in points]
            found.append(dict(format=getattr(result.format,'name',str(result.format)),box=[max(0,min(1,min(xs)/width)),max(0,min(1,min(ys)/height)),max(0,min(1,max(xs)/width)),max(0,min(1,max(ys)/height))]))
        return found, None
    except ImportError:
        return [], 'Local barcode scanner is unavailable; code presence was not verified.'
    except Exception:
        return [], 'Local barcode scanner could not read this rendered page; code presence remains unverified.'

def extract(path, role, outdir, artifact_id, name):
    from PIL import Image, ImageOps
    p = Path(path); out = Path(outdir); out.mkdir(parents=True, exist_ok=True)
    doc = dict(id=artifact_id, name=name, role=role, sha256=hashlib.sha256(p.read_bytes()).hexdigest(),
               bytes=p.stat().st_size, pages=[], cells=[], lines=[], warnings=[], text='', data=None)
    ext = p.suffix.lower()
    if ext in ('.xlsx', '.docx'):
        with zipfile.ZipFile(p) as archive:
            if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 80_000_000:
                raise ValueError('Expanded Office document exceeds the safe parsing limit.')
    if ext == '.pdf':
        import pypdfium2 as pdfium
        if not p.read_bytes()[:1024].lstrip().startswith(b'%PDF-'):
            raise ValueError('File does not have a PDF signature.')
        pdf = pdfium.PdfDocument(str(p))
        if len(pdf) > MAX_PAGES: raise ValueError('Limit: 12 PDF pages per file. Split the document into relevant sections.')
        for n in range(len(pdf)):
            page = pdf[n]; width, height = page.get_size(); left, bottom, right, top = page.get_bbox()
            if width <= 0 or height <= 0: raise ValueError('Invalid PDF page dimensions.')
            image_name = f'{artifact_id}-p{n+1}.png'
            rendered=page.render(scale=min(2, 1800/max(width,height))).to_pil().convert('RGB')
            rendered.save(out/image_name)
            codes, scan_warning=scan_codes(rendered)
            tp = page.get_textpage(); text = tp.get_text_range().replace('\x00','')
            if len(text) > 100_000: raise ValueError('Too much text in one PDF page.')
            lines=[]
            for m in re.finditer(r'[^\r\n]+', text):
                value=m.group().strip()
                if not value: continue
                boxes=[]
                for ix in range(m.start(), min(m.end(),tp.count_chars())):
                    try:
                        b=tp.get_charbox(ix)
                        if b[2]>b[0] and b[3]>b[1]: boxes.append(b)
                    except Exception: pass
                rect=None
                if boxes and page.get_rotation()==0:
                    x0=min(b[0] for b in boxes); y0=min(b[1] for b in boxes)
                    x1=max(b[2] for b in boxes); y1=max(b[3] for b in boxes)
                    rect=[max(0,min(1,(x0-left)/width)), max(0,min(1,(top-y1)/height)),
                          max(0,min(1,(x1-left)/width)), max(0,min(1,(top-y0)/height))]
                lines.append(dict(text=value,page=n+1,box=rect))
            doc['lines'] += lines
            doc['pages'].append(dict(number=n+1,width=width,height=height,unit='pt',image=image_name,text=text,codes=codes,
                                      method='pdf_text' if len(text.strip())>40 else 'vision_needed'))
            if scan_warning and scan_warning not in doc['warnings']: doc['warnings'].append(scan_warning)
            if len(text.strip())<40: doc['warnings'].append(f'Page {n+1} has little extractable text. Vision or manual review is needed.')
            tp.close(); page.close()
        pdf.close()
        doc['text']='\n\n'.join(x['text'] for x in doc['pages'])
    elif ext in ('.png','.jpg','.jpeg','.webp'):
        Image.MAX_IMAGE_PIXELS=40_000_000
        im=Image.open(p); im=ImageOps.exif_transpose(im).convert('RGB')
        width,height=im.size
        if width*height>40_000_000: raise ValueError('Image exceeds 40 megapixels.')
        im.thumbnail((1800,1800)); image_name=f'{artifact_id}-p1.png'; im.save(out/image_name)
        codes, scan_warning=scan_codes(im)
        doc['pages']=[dict(number=1,width=width,height=height,unit='px',image=image_name,text='',codes=codes,method='vision_needed')]
        if scan_warning: doc['warnings'].append(scan_warning)
        doc['warnings']=['Image input requires connected vision analysis or manual review; no local OCR text was inferred.']
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True,data_only=True,keep_links=False)
        for sheet in wb:
            if sheet.max_row>2000 or sheet.max_column>80: raise ValueError('Spreadsheet limit: 2,000 rows and 80 columns per sheet.')
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        value=str(cell.value)
                        doc['cells'].append(dict(sheet=sheet.title,cell=cell.coordinate,row=cell.row,col=cell.column,value=value))
        wb.close()
        doc['text']='\n'.join(f"{c['sheet']}!{c['cell']}: {c['value']}" for c in doc['cells'])
        doc['warnings'].append('Spreadsheet formulas use their saved values; formulas, macros, and external links are not executed.')
    elif ext=='.docx':
        from docx import Document
        d=Document(p)
        doc['lines']=[dict(text=x.text,paragraph=i+1) for i,x in enumerate(d.paragraphs) if x.text.strip()]
        for ti,table in enumerate(d.tables):
            for ri,row in enumerate(table.rows):
                for ci,cell in enumerate(row.cells):
                    if cell.text.strip(): doc['cells'].append(dict(sheet=f'Table {ti+1}',cell=f'R{ri+1}C{ci+1}',row=ri+1,col=ci+1,value=cell.text))
        doc['text']='\n'.join(x['text'] for x in doc['lines'])+'\n'+'\n'.join(f"{x['sheet']}!{x['cell']}: {x['value']}" for x in doc['cells'])
        doc['warnings'].append('Embedded images in Word files are not analyzed. Export image-heavy documents as PDF.')
    elif ext in ('.json','.csv','.txt','.md'):
        raw=p.read_text(encoding='utf-8-sig')
        if len(raw)>150_000: raise ValueError('Text input exceeds 150,000 characters.')
        doc['text']=raw
        doc['lines']=[dict(text=line,line=i+1) for i,line in enumerate(raw.splitlines()) if line.strip()]
        if ext=='.json': doc['data']=json.loads(raw)
        if ext=='.csv':
            rows=list(csv.reader(io.StringIO(raw)))
            if len(rows)>2000: raise ValueError('CSV limit is 2,000 rows.')
            doc['rows']=rows
    else: raise ValueError('Supported formats: PDF, PNG, JPEG, WebP, XLSX, DOCX, CSV, JSON, TXT and MD.')
    if len(doc['text'])>200_000: raise ValueError('Document text exceeds 200,000 characters; upload relevant sections only.')
    return doc

def report(review_path, output):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, Image as ReportImage, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    r=json.loads(Path(review_path).read_text()); styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='PP',fontName='Helvetica',fontSize=9,leading=13,spaceAfter=7,splitLongWords=True))
    styles.add(ParagraphStyle(name='PPH',fontName='Helvetica-Bold',fontSize=12,leading=16,spaceBefore=12,spaceAfter=6,textColor=HexColor('#175b48'),keepWithNext=True))
    def p(s,style='PP'):
        s=str(s).replace('→',' -> ').replace('↔',' / ').replace('✓','match').replace('–','-').replace('—','-')
        return Paragraph(escape(s).replace('\n','<br/>'),styles[style])
    story=[p('PACKSIGNAL | Packaging review','Title'),Spacer(1,14),p(r['name'],'Heading1'),p(f"Review {r['id']} | {r['createdAt']}"),p(f"Market: {r.get('market','Unspecified')} | Engine: {r.get('engine','local')}"),p('Scope: comparison against supplied sources. This report does not certify legal compliance or authorize printing.'),p('Results','PPH')]
    counts=r.get('counts',{})
    story.append(p(' | '.join(f'{k}: {v}' for k,v in counts.items())))
    for w in r.get('warnings',[]):story.append(p('Review note: '+w))
    for limit in r.get('coverage',{}).get('limits',[]):story.append(p('Coverage: '+limit))
    for f in r.get('findings',[]):
        finding_story=[p(f"{f['status'].upper()} | {f['title']}",'PPH'),p('Expected: '+str(f.get('expected',''))),p('Observed: '+str(f.get('observed',''))),p(f.get('explanation',''))]
        if f.get('rule'):finding_story.append(p('Rule: '+f['rule'].get('text','')))
        seen_evidence=set()
        for ev in f.get('evidence',[]):
            evidence_key=json.dumps({key:ev.get(key) for key in ('documentId','name','location','page','box','quote')},sort_keys=True,ensure_ascii=False)
            if evidence_key in seen_evidence:continue
            seen_evidence.add(evidence_key)
            finding_story.append(p(f"Evidence - {ev.get('name','')} {ev.get('location','')}: {ev.get('quote','')}"))
        decision=r.get('decisions',{}).get(f['id'])
        if decision:finding_story.append(p(f"Reviewer: {decision['disposition']} - {decision.get('note','')} ({decision['at']})"+(' [Needs review: source comparison changed]' if decision.get('stale') else '')))
        story.append(KeepTogether(finding_story+[Spacer(1,2)]))
    if r.get('revision'):
        story.append(p('Revision comparison','PPH'))
        story.append(p('Previous review: '+r['revision']['previousId']))
        for change in r['revision']['changes']:story.append(p(f"{change['title']}: {change.get('before') or 'new'} -> {change.get('after') or 'not rechecked'} ({change['change']})"))
    story.append(p('Source inventory','PPH'))
    for d in r.get('documents',[]):story.append(p(f"{d['role']}: {d['name']} | SHA-256 {d['sha256']}"))
    if r.get('signature'):
        signature=r['signature']
        story.extend([p('Cryptographic review seal','PPH'),p(f"{signature.get('algorithm','Signature')} | key {signature.get('publicKeyId','unknown')} | sealed {signature.get('sealedAt','')}"),p('Verification hash: '+str(signature.get('payloadHash',''))),p('Signature (Base64): '+str(signature.get('signature',''))),p('Public verification key: '+str(signature.get('publicKey',''))),p('The signed payload and signature are also available in the companion HTML or JSON export. Trust the public-key fingerprint through a separate channel to establish signer identity.')])
    story.append(p('Audit history','PPH'))
    for e in r.get('audit',[]):
        story.append(p(f"{e['at']} | {e['action']} | {e.get('detail','')}"))
        if e.get('options'):story.append(p('Recorded choices: '+json.dumps(e['options'])))
    art=next((d for d in r.get('documents',[]) if d['role']=='artwork'),None)
    if art:
        story.extend([PageBreak(),p('Artwork reference','Heading1'),p('Reduced overview of supplied pages. Use the application and original source for detailed visual inspection.')])
        for page in art['pages']:
            image=Path(review_path).parent/'rendered'/page['image']
            if image.exists():
                width=490; height=width*page['height']/page['width']
                if height>600:width*=600/height;height=600
                story.extend([p(f"{art['name']} - page {page['number']} ({page['method']})"),ReportImage(str(image),width=width,height=height),Spacer(1,12)])
    def footer(c,d):
        c.setFont('Helvetica',8);c.setFillColor(HexColor('#68766f'));c.drawString(42,28,'PackSignal - human release decision required')
        if r.get('signature'):
            c.drawCentredString(A4[0]/2,12,'Seal '+str(r['signature'].get('payloadHash',''))[:16])
        c.drawRightString(A4[0]-42,28,str(d.page))
    SimpleDocTemplate(output,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=52).build(story,onFirstPage=footer,onLaterPages=footer)
    return {'path':output}

if __name__=='__main__':
    try:
        args=json.load(sys.stdin)
        result=report(args['review'],args['output']) if args.get('action')=='report' else extract(args['path'],args['role'],args['outdir'],args['id'],args['name'])
        print(json.dumps(result,ensure_ascii=False))
    except Exception as e:
        print(json.dumps({'error':str(e)[:500]}));sys.exit(1)
